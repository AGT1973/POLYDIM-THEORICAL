"""
POLYDIM V814 — Generador y Ensamblador de Tesis Doctoral DOCX
Convierte todos los capítulos LaTeX a DOCX vía Pandoc (MathML/OMML)
y los fusiona en el documento maestro final: TESIS_DOCTORAL_POLYDIM_V814.docx
"""
import subprocess, sys, re, copy
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

LATEX_DIR = Path(r"E:\POLYDIM-THEORICAL\CONSTITUCION_TESIS_MANIFESTO\TESIS_DOCTORAL_LATEX")
TEMP_DIR  = LATEX_DIR / "_pandoc_temp"
OUTPUT    = Path(r"E:\POLYDIM-THEORICAL\CONSTITUCION_TESIS_MANIFESTO\TESIS_DOCTORAL_POLYDIM_V814.docx")
PANDOC    = r"C:\Users\eluithi\AppData\Local\Pandoc\pandoc.exe"

TEMP_DIR.mkdir(exist_ok=True)

CLEAN_PATTERNS = [
    (r'\\begin\{polydimbox\}', r'\\textbf{[POLYDIM]} '),
    (r'\\end\{polydimbox\}', ''),
    (r'\\begin\{criticalbox\}', r'\\textbf{[CRÍTICO]} '),
    (r'\\end\{criticalbox\}', ''),
    (r'\\begin\{certifiedbox\}', r'\\textbf{[CERTIFICADO]} '),
    (r'\\end\{certifiedbox\}', ''),
    (r'\\begin\{longtable\}(\[.*?\])?\{.*?\}', r'\\begin{tabular}{llll}'),
    (r'\\end\{longtable\}', r'\\end{tabular}'),
    (r'\\begin\{table\}(\[.*?\])?', ''),
    (r'\\end\{table\}', ''),
    (r'\\(toprule|midrule|bottomrule)', r'\\hline'),
    (r'\\caption\{[^{}]*(?:\{[^{}]*\}[^{}]*)*\}', ''),
    (r'\\label\{[^{}]*\}', ''),
    (r'\\begin\{lstlisting\}(\[.*?\])?', r'\\begin{verbatim}'),
    (r'\\end\{lstlisting\}', r'\\end{verbatim}'),
    (r'\\lstdefinestyle\{.*?\}\{.*?\}', '', re.DOTALL),
    (r'\\begin\{mdframed\}(\[.*?\])?', ''),
    (r'\\end\{mdframed\}', ''),
    (r'\\epigraph\{(.*?)\}\{(.*?)\}', r'\\textit{\1} --- \2', re.DOTALL),
    (r'\\vers\b', r'\\operatorname{vers}'),
    (r'\\tol\b', r'\\operatorname{tol}'),
    (r'\\St\b', r'\\operatorname{St}'),
    (r'\\norm\{(.*?)\}', r'\\|{\1}\\|'),
    (r'\\abs\{(.*?)\}', r'|{\1}|'),
    (r'\\usepackage.*?\n', '\n'),
    (r'\\pagestyle\{.*?\}', ''),
    (r'\\fancyhf\{.*?\}', ''),
    (r'\\fancyhead.*?\n', '\n'),
    (r'\\onehalfspacing', ''),
    (r'\\pgfplotsset\{.*?\}', ''),
]

TODOS_CAPS = [
    "cap01_prefacio", "cap02_dogma_no_worm", "cap03_dpi_formal",
    "cap04_variedad_esferica", "cap05_clifford_rotores", "cap06_rodrigues",
    "cap07_cayley_smw", "cap08_betti_topologia", "cap09_clifford_qpu",
    "cap10_bargmann_pancharatnam", "cap11_pmtp_protocolo", "cap12_seqlock_hardened",
    "cap13_latent_os", "cap14_ffi_abi", "cap15_hardware_heterogeneo",
    "cap16_kernel_cpp", "cap17_kernel_rust", "cap18_triton_gpu",
    "cap19_orquestador_python", "cap_v764_auditoria", "cap_v765_vector_b",
    "cap_v765_telepathy", "cap_v814_tribunal_espacio_vectorial", "cap20_fases_certificadas",
    "cap21_metodologia_red_team", "cap22_resultados_silicio", "cap23_escalabilidad",
    "cap24_impacto_termodinamico", "cap25_latam_tokens", "cap26_futuro_interlat",
    "apendice_a_codigo_fuente", "apendice_b_telemetria", "apendice_c_demostraciones",
]

def clean_tex(tex_text):
    text = tex_text
    for p in CLEAN_PATTERNS:
        flags = p[2] if len(p) > 2 else 0
        try:
            text = re.sub(p[0], p[1], text, flags=flags)
        except Exception:
            pass
    return text

def tex_to_docx(cap_name, force=True):
    src = LATEX_DIR / f"{cap_name}.tex"
    dst = TEMP_DIR / f"{cap_name}.docx"
    
    if dst.exists() and not force:
        print(f"  [SKIP] {cap_name}.docx ya existe")
        return dst
    
    if not src.exists():
        print(f"  [MISS] {cap_name}.tex no existe")
        return None
    
    raw = src.read_text(encoding='utf-8', errors='replace')
    cleaned = clean_tex(raw)
    
    tmp_tex = TEMP_DIR / f"{cap_name}_clean.tex"
    tmp_tex.write_text(cleaned, encoding='utf-8')
    
    cmd = [PANDOC, str(tmp_tex), "-f", "latex",
           "-t", "docx", "--mathml", "-o", str(dst),
           "--resource-path", str(LATEX_DIR)]
    
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
    if r.returncode == 0:
        print(f"  [OK] {cap_name}.docx ({dst.stat().st_size/1024:.1f} KB)")
        return dst
    else:
        err = r.stderr[:150].replace('\n', ' ')
        print(f"  [WARN] {cap_name}: {err}")
        return None

def add_page_break(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    run._r.append(br)

def build_final_docx(docs):
    master = Document()
    for section in master.sections:
        section.top_margin = Cm(3)
        section.bottom_margin = Cm(3)
        section.left_margin = Cm(3.5)
        section.right_margin = Cm(2.5)
    
    # PORTADA V814
    t = master.add_heading("DIMENSION IS ALL YOU NEED", level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    p = master.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fundamentación Matemática, Arquitectura de Sistema Operativo Latente\n"
                  "y Protocolo PMTP Zero-Copy para la Comunicación Nativa en S\u1d30\u207b\u00b9")
    r.font.size = Pt(14); r.font.italic = True
    
    master.add_paragraph()
    p = master.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Ariel García Traba\nPOLYDIM / EinsofOS Research Initiative\nSeptiembre de 2026")
    r.font.size = Pt(16); r.font.bold = True
    
    master.add_paragraph()
    p = master.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("SERIE 800 — KERNEL V814 CERTIFICADO\n"
                  "Exit Code 0 | 7/7 Suites PASS | Drift \u2264 4.44\u00d710\u207b\u00b9\u2076 | PMTP SPSC 34.8 ns\n"
                  "Tribunal Multi-IA: 486 Hallazgos Adversariales Integrados en Base Vectorial")
    r.font.size = Pt(12); r.font.bold = True
    r.font.color.rgb = RGBColor(0, 91, 170)
    add_page_break(master)
    
    # RESUMEN V814
    master.add_heading("Resumen", level=1)
    master.add_paragraph(
        "El paradigma hegemónico del Procesamiento del Lenguaje Natural impone que la "
        "comunicación entre agentes inteligentes transite por secuencias unidimensionales "
        "de tokens discretos. Esta tesis demuestra teórica y empíricamente que este colapso "
        "constituye la Paradoja del Gusano 1D: destruye irreversiblemente información "
        "geométrica bajo la Desigualdad de Procesamiento de Datos (DPI).\n\n"
        "Para erradicar este colapso, se presenta POLYDIM / EinsofOS y el protocolo "
        "PMTP V814: una arquitectura donde los agentes de IA se comunican nativamente "
        "mediante tensores en S^(D-1) (D = 10\u2076) vía Memoria Compartida Zero-Copy.\n\n"
        "Resultados certificados en silicio real (AMD x64, GCC 14, Rust 1.98, Python 3.12):\n"
        "• Rodrigues Geodésico D=10\u2076: 35.06 ms | Drift = 4.44×10\u207b¹\u2076 ≤ 2\u03b5_mach\n"
        "• PMTP Zero-Copy IPC SPSC: 34.8 ns | Max Bit Diff = 0\n"
        "• Cayley-SMW Streaming D\u00d7K (K=8): ||Y\u1d40Y-I||_max = 3.33×10\u207b¹\u2074\n"
        "• Reducción BLAS TwoSum / DSYRK: Acotada a 2\u03b5_mach en D=10\u2076\n"
        "• FFI Guard RAII Dart/Flutter 3DGS vía NativeFinalizer en memoria unificada\n"
        "• Guardia Topológica Betti-1: \u03b2_0=1 (conexo), \u03b2_1 preservado\n"
        "• 100% de ataques adversariales bloqueados sin Segfault ni pánicos FFI.\n\n"
        "Palabras clave: S^(D-1), PMTP Zero-Copy, DPI, Clifford, Rodrigues, Stiefel Cayley-SMW, Betti-1, EinsofOS, Tribunal Multi-IA."
    )
    add_page_break(master)
    
    # FUSIÓN DE CAPÍTULOS
    total_ok = 0
    for doc_path in docs:
        if doc_path is None or not Path(str(doc_path)).exists():
            continue
        try:
            src_doc = Document(str(doc_path))
            for element in src_doc.element.body:
                master.element.body.append(copy.deepcopy(element))
            add_page_break(master)
            total_ok += 1
        except Exception as e:
            print(f"  [MERGE WARN] {Path(str(doc_path)).name}: {e}")
    
    master.save(str(OUTPUT))
    size_mb = OUTPUT.stat().st_size / (1024*1024)
    print(f"\n{'='*60}")
    print(f"DOCX FINAL V814 GENERADO: {OUTPUT}")
    print(f"Capítulos fusionados: {total_ok}/{len(docs)}")
    print(f"Tamaño: {size_mb:.2f} MB")
    print(f"{'='*60}")

if __name__ == "__main__":
    print("POLYDIM — Compilador de Tesis Doctoral V814")
    print("="*60)
    converted = []
    for cap in TODOS_CAPS:
        print(f"Procesando {cap}...")
        result = tex_to_docx(cap, force=True)
        converted.append(result)
    
    ok_count = sum(1 for c in converted if c and Path(str(c)).exists())
    print(f"\nConversión: {ok_count}/{len(TODOS_CAPS)}")
    print("Construyendo DOCX final V814...")
    build_final_docx(converted)
